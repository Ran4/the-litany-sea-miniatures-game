# /// script
# requires-python = ">=3.11"
# dependencies = ["python-docx"]
# ///
"""
SIXES rulebook generator (data-driven).
Reads sixes_content.json (produced by the worldbuilding workflow) and the static
SIXES core engine, and emits sixes_rulebook.docx.

Run:  uv run sixes_book.py
"""
import json
import os
from docx import Document
from docx.shared import Pt, Mm, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

# ---- palette -----------------------------------------------------------------
INK     = RGBColor(0x20, 0x20, 0x20)
ACCENT  = RGBColor(0x7A, 0x14, 0x14)   # deep blood red
STEEL   = RGBColor(0x2E, 0x45, 0x52)   # slate
GREY    = RGBColor(0x5A, 0x5A, 0x5A)
WHITE   = RGBColor(0xFF, 0xFF, 0xFF)
GOLD    = RGBColor(0x8A, 0x6D, 0x1F)
HDR_FILL   = "2E4552"
ZEBRA_FILL = "ECEFF1"
BOX_FILL   = "F4EFE6"
CHAR_FILL  = "ECE6F0"
PART_FILL  = "2E4552"
BODY_FONT = "Calibri"

# ---- load data ---------------------------------------------------------------
HERE = os.path.dirname(os.path.abspath(__file__))
with open(os.path.join(HERE, "sixes_content.json"), encoding="utf-8") as fh:
    DATA = json.load(fh)

setting  = DATA.get("setting", {}) or {}
prose    = DATA.get("prose", {}) or {}
council  = DATA.get("council", {}) or {}
factions = DATA.get("factions", []) or []
scenarios = DATA.get("scenarios", []) or []
campaign = DATA.get("campaign", {}) or {}
advanced = DATA.get("advanced", {}) or {}

WORLD = setting.get("worldName", "SIXES")
TAGLINE = setting.get("tagline", "A Fast Skirmish Wargame")

_epi_pool = list(prose.get("sectionEpigraphs", []) or [])
_epi_idx = 0


def next_epigraph():
    global _epi_idx
    if not _epi_pool:
        return None
    e = _epi_pool[_epi_idx % len(_epi_pool)]
    _epi_idx += 1
    return e


ASSETS = os.path.join(HERE, "assets")
IMG_CAP = {}
_ipath = os.path.join(HERE, "images_prompts.json")
if os.path.exists(_ipath):
    try:
        for _im in json.load(open(_ipath, encoding="utf-8")):
            IMG_CAP[_im.get("id")] = _im.get("caption", "")
    except Exception:
        pass


def has_img(iid):
    p = os.path.join(ASSETS, iid + ".png")
    return p if (os.path.exists(p) and os.path.getsize(p) > 5000) else None


# ---- low-level helpers -------------------------------------------------------
def shade(cell, hex_fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_fill)
    tcPr.append(shd)


def cell_text(cell, text, bold=False, color=None, align=None, size=None):
    cell.text = ""
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(1)
    p.paragraph_format.space_after = Pt(1)
    if align is not None:
        p.alignment = align
    run = p.add_run("" if text is None else str(text))
    run.font.name = BODY_FONT
    run.bold = bold
    if size:
        run.font.size = Pt(size)
    if color is not None:
        run.font.color.rgb = color
    return run


def set_col_widths(table, widths):
    table.autofit = False
    for row in table.rows:
        for i, w in enumerate(widths):
            if i < len(row.cells):
                row.cells[i].width = w


def add_table(doc, headers, rows, widths=None, font_size=10):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, h in enumerate(headers):
        c = table.rows[0].cells[i]
        shade(c, HDR_FILL)
        cell_text(c, h, bold=True, color=WHITE, size=font_size, align=WD_ALIGN_PARAGRAPH.CENTER)
    for r_idx, row in enumerate(rows):
        cells = table.add_row().cells
        for i, val in enumerate(row):
            if i >= len(cells):
                continue
            align = WD_ALIGN_PARAGRAPH.LEFT if i == 0 else WD_ALIGN_PARAGRAPH.CENTER
            cell_text(cells[i], val, size=font_size, align=align)
            if r_idx % 2 == 1:
                shade(cells[i], ZEBRA_FILL)
    if widths:
        set_col_widths(table, widths)
    doc.add_paragraph().paragraph_format.space_after = Pt(3)
    return table


def _runs_from_markup(p, text, size, color=INK):
    for i, seg in enumerate(str(text).split("**")):
        if seg == "":
            continue
        run = p.add_run(seg)
        run.font.name = BODY_FONT
        run.font.size = Pt(size)
        run.font.color.rgb = color
        if i % 2 == 1:
            run.bold = True


def heading(doc, text, level=1):
    p = doc.add_paragraph()
    if level == 1:
        p.paragraph_format.space_before = Pt(14)
        p.paragraph_format.space_after = Pt(4)
        run = p.add_run(text.upper())
        run.font.size = Pt(16)
        run.font.bold = True
        run.font.color.rgb = ACCENT
        pPr = p._p.get_or_add_pPr()
        pbdr = OxmlElement("w:pBdr")
        bottom = OxmlElement("w:bottom")
        bottom.set(qn("w:val"), "single")
        bottom.set(qn("w:sz"), "12")
        bottom.set(qn("w:space"), "3")
        bottom.set(qn("w:color"), "7A1414")
        pbdr.append(bottom)
        pPr.append(pbdr)
    else:
        p.paragraph_format.space_before = Pt(8)
        p.paragraph_format.space_after = Pt(2)
        run = p.add_run(text)
        run.font.size = Pt(12.5)
        run.font.bold = True
        run.font.color.rgb = STEEL
    run.font.name = BODY_FONT
    return p


def body(doc, text, size=10.5, space_after=6, italic=False, align=None, color=INK):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing = 1.08
    if align is not None:
        p.alignment = align
    parts = str(text).split("**")
    for i, seg in enumerate(parts):
        if seg == "":
            continue
        run = p.add_run(seg)
        run.font.name = BODY_FONT
        run.font.size = Pt(size)
        run.font.color.rgb = color
        run.italic = italic
        if i % 2 == 1:
            run.bold = True
    return p


def bullet(doc, text, size=10.5):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.space_after = Pt(2)
    _runs_from_markup(p, text, size)
    return p


def numbered(doc, text, size=10.5):
    p = doc.add_paragraph(style="List Number")
    p.paragraph_format.space_after = Pt(2)
    _runs_from_markup(p, text, size)
    return p


def _accent_box(doc, fill, accent_hex):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = table.rows[0].cells[0]
    shade(cell, fill)
    cell.width = Mm(170)
    tcPr = cell._tc.get_or_add_tcPr()
    borders = OxmlElement("w:tcBorders")
    left = OxmlElement("w:left")
    left.set(qn("w:val"), "single")
    left.set(qn("w:sz"), "26")
    left.set(qn("w:space"), "0")
    left.set(qn("w:color"), accent_hex)
    borders.append(left)
    tcPr.append(borders)
    cell.text = ""
    return cell


def callout(doc, title, text, fill=BOX_FILL, accent_hex="7A1414"):
    cell = _accent_box(doc, fill, accent_hex)
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(2)
    r = p.add_run(title.upper())
    r.bold = True
    r.font.size = Pt(10)
    r.font.name = BODY_FONT
    r.font.color.rgb = ACCENT
    p2 = cell.add_paragraph()
    p2.paragraph_format.space_after = Pt(2)
    _runs_from_markup(p2, text, 10)
    doc.add_paragraph().paragraph_format.space_after = Pt(3)


def epigraph(doc, e):
    if not e:
        return
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.left_indent = Mm(8)
    r = p.add_run('“' + e.get("quote", "") + '”')
    r.italic = True
    r.font.size = Pt(10.5)
    r.font.name = BODY_FONT
    r.font.color.rgb = STEEL
    p2 = doc.add_paragraph()
    p2.paragraph_format.space_after = Pt(8)
    p2.paragraph_format.left_indent = Mm(8)
    r2 = p2.add_run("— " + e.get("attribution", ""))
    r2.italic = True
    r2.font.size = Pt(9.5)
    r2.font.name = BODY_FONT
    r2.font.color.rgb = GREY


def fiction(doc, text):
    for para in [pp for pp in str(text).split("\n") if pp.strip()]:
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(6)
        p.paragraph_format.line_spacing = 1.12
        r = p.add_run(para.strip())
        r.italic = True
        r.font.size = Pt(11)
        r.font.name = BODY_FONT
        r.font.color.rgb = INK


def add_image(doc, iid, width_in, caption=None, center=True):
    p = has_img(iid)
    if not p:
        return False
    par = doc.add_paragraph()
    if center:
        par.alignment = WD_ALIGN_PARAGRAPH.CENTER
    par.paragraph_format.space_before = Pt(4)
    par.add_run().add_picture(p, width=Inches(width_in))
    cap = caption if caption is not None else IMG_CAP.get(iid, "")
    if cap:
        cp = doc.add_paragraph()
        cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        cp.paragraph_format.space_after = Pt(8)
        cr = cp.add_run(cap)
        cr.italic = True
        cr.font.size = Pt(8.5)
        cr.font.color.rgb = GREY
        cr.font.name = BODY_FONT
    return True


def add_image_path(doc, path, width_in, caption=None, center=True):
    if not (os.path.exists(path) and os.path.getsize(path) > 5000):
        return False
    par = doc.add_paragraph()
    if center:
        par.alignment = WD_ALIGN_PARAGRAPH.CENTER
    par.paragraph_format.space_before = Pt(4)
    par.add_run().add_picture(path, width=Inches(width_in))
    if caption:
        cp = doc.add_paragraph()
        cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        cp.paragraph_format.space_after = Pt(8)
        cr = cp.add_run(caption)
        cr.italic = True
        cr.font.size = Pt(8.5)
        cr.font.color.rgb = GREY
        cr.font.name = BODY_FONT
    return True


def two_up_images(doc, left_id, right_id, captions):
    lp, rp = has_img(left_id), has_img(right_id)
    if not (lp or rp):
        return
    table = doc.add_table(rows=1, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for idx, (iid, pth) in enumerate([(left_id, lp), (right_id, rp)]):
        cell = table.rows[0].cells[idx]
        cell.width = Mm(83)
        cell.text = ""
        cpar = cell.paragraphs[0]
        cpar.alignment = WD_ALIGN_PARAGRAPH.CENTER
        if pth:
            cpar.add_run().add_picture(pth, width=Inches(2.95))
        capp = cell.add_paragraph()
        capp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        cr = capp.add_run(captions.get(iid, IMG_CAP.get(iid, "")))
        cr.italic = True
        cr.font.size = Pt(8)
        cr.font.color.rgb = GREY
        cr.font.name = BODY_FONT
    doc.add_paragraph().paragraph_format.space_after = Pt(4)


def part_divider(doc, label, title, blurb=None):
    doc.add_page_break()
    for _ in range(3):
        doc.add_paragraph()
    p0 = doc.add_paragraph()
    p0.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r0 = p0.add_run(label.upper())
    r0.font.size = Pt(13)
    r0.font.bold = True
    r0.font.name = BODY_FONT
    r0.font.color.rgb = GREY
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(title.upper())
    r.font.size = Pt(34)
    r.font.bold = True
    r.font.name = BODY_FONT
    r.font.color.rgb = ACCENT
    if blurb:
        pb = doc.add_paragraph()
        pb.alignment = WD_ALIGN_PARAGRAPH.CENTER
        pb.paragraph_format.space_before = Pt(6)
        rb = pb.add_run(blurb)
        rb.italic = True
        rb.font.size = Pt(12)
        rb.font.name = BODY_FONT
        rb.font.color.rgb = STEEL
    doc.add_page_break()


def statline(d):
    def g(k):
        v = d.get(k, "—")
        return str(v) if v not in (None, "") else "—"
    return f"M {g('M')}  ·  SK {g('SK')}  ·  DEF {g('DEF')}  ·  W {g('W')}  ·  A {g('A')}"


def character_box(doc, ch):
    cell = _accent_box(doc, CHAR_FILL, "5A3E78")
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(1)
    r = p.add_run(ch.get("name", "Unknown"))
    r.bold = True
    r.font.size = Pt(12)
    r.font.name = BODY_FONT
    r.font.color.rgb = RGBColor(0x4A, 0x2E, 0x68)
    title = ch.get("title")
    if title:
        rt = p.add_run("  —  " + title)
        rt.italic = True
        rt.font.size = Pt(10)
        rt.font.name = BODY_FONT
        rt.font.color.rgb = GREY
    ps = cell.add_paragraph()
    ps.paragraph_format.space_after = Pt(1)
    rs = ps.add_run(statline(ch))
    rs.font.size = Pt(9.5)
    rs.bold = True
    rs.font.name = BODY_FONT
    rs.font.color.rgb = STEEL
    pts = ch.get("points")
    if pts is not None:
        rp = ps.add_run(f"   —   {pts} pts")
        rp.font.size = Pt(9.5)
        rp.font.name = BODY_FONT
        rp.font.color.rgb = GREY
    weapons = ch.get("weapons") or []
    if weapons:
        pw = cell.add_paragraph()
        pw.paragraph_format.space_after = Pt(1)
        _runs_from_markup(pw, "**Wargear:** " + ", ".join(weapons), 9.5)
    ur = ch.get("uniqueRule") or {}
    if ur:
        pu = cell.add_paragraph()
        pu.paragraph_format.space_after = Pt(1)
        _runs_from_markup(pu, "**" + ur.get("name", "Rule") + ":** " + ur.get("text", ""), 9.5)
    flv = ch.get("flavor")
    if flv:
        pf = cell.add_paragraph()
        pf.paragraph_format.space_after = Pt(1)
        rf = pf.add_run(flv)
        rf.italic = True
        rf.font.size = Pt(9.5)
        rf.font.name = BODY_FONT
        rf.font.color.rgb = INK
    doc.add_paragraph().paragraph_format.space_after = Pt(3)


# ---- document setup ----------------------------------------------------------
doc = Document()
normal = doc.styles["Normal"]
normal.font.name = BODY_FONT
normal.font.size = Pt(10.5)
normal.font.color.rgb = INK

sec = doc.sections[0]
sec.page_height = Mm(297)
sec.page_width = Mm(210)
sec.top_margin = Inches(0.7)
sec.bottom_margin = Inches(0.65)
sec.left_margin = Inches(0.85)
sec.right_margin = Inches(0.85)
fp = sec.footer.paragraphs[0]
fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
fr = fp.add_run(f"{WORLD}  ·  a SIXES game  ·  v1.0")
fr.font.size = Pt(8)
fr.font.color.rgb = GREY
fr.font.name = BODY_FONT


# ============================================================ TITLE PAGE
for _ in range(4):
    doc.add_paragraph()
t = doc.add_paragraph()
t.alignment = WD_ALIGN_PARAGRAPH.CENTER
tr = t.add_run(WORLD.upper())
tr.font.name = BODY_FONT
tr.font.size = Pt(60)
tr.font.bold = True
tr.font.color.rgb = ACCENT

sub = doc.add_paragraph()
sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
sr = sub.add_run(TAGLINE)
sr.italic = True
sr.font.size = Pt(15)
sr.font.color.rgb = STEEL
sr.font.name = BODY_FONT

doc.add_paragraph()
ps = doc.add_paragraph()
ps.alignment = WD_ALIGN_PARAGRAPH.CENTER
rps = ps.add_run("Powered by the SIXES engine")
rps.font.size = Pt(12)
rps.font.color.rgb = GREY
rps.font.name = BODY_FONT

_cover = has_img("cover")
if _cover:
    cv = doc.add_paragraph()
    cv.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cv.paragraph_format.space_before = Pt(8)
    cv.paragraph_format.space_after = Pt(8)
    cv.add_run().add_picture(_cover, width=Inches(6.2))
else:
    for _ in range(8):
        doc.add_paragraph()
need = doc.add_paragraph()
need.alignment = WD_ALIGN_PARAGRAPH.CENTER
nr = need.add_run("ALL YOU NEED")
nr.bold = True
nr.font.size = Pt(11)
nr.font.color.rgb = ACCENT
nr.font.name = BODY_FONT
for line in ["A handful of six-sided dice  ·  A tape measure",
             "Terrain  ·  Two warbands  ·  A worthy opponent"]:
    pp = doc.add_paragraph()
    pp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    rr = pp.add_run(line)
    rr.font.size = Pt(11)
    rr.font.color.rgb = INK
    rr.font.name = BODY_FONT
for _ in range(6):
    doc.add_paragraph()
ver = doc.add_paragraph()
ver.alignment = WD_ALIGN_PARAGRAPH.CENTER
vr = ver.add_run("The Complete Rulebook  ·  Version 1.0")
vr.font.size = Pt(10)
vr.font.color.rgb = GREY
vr.font.name = BODY_FONT
doc.add_page_break()


# ============================================================ OPENING FICTION
if prose.get("openingFiction"):
    for _ in range(2):
        doc.add_paragraph()
    fiction(doc, prose["openingFiction"])
    doc.add_page_break()


# ============================================================ CONTENTS
heading(doc, "Contents", 1)
toc = [
    "Part I  —  The World",
    "Part II  —  Core Rules (the SIXES engine)",
    "Part III  —  The Factions",
]
for f in factions:
    toc.append("        · " + f.get("name", "Faction"))
toc += [
    "Part IV  —  Scenarios",
    "Part V  —  Campaign: " + (campaign.get("name", "The Long War")),
    "Part VI  —  Advanced & Optional Rules",
    "Part VII  —  Quick Reference",
]
for c in toc:
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(3)
    r = p.add_run(c)
    r.font.size = Pt(11.5)
    r.font.color.rgb = INK
    r.font.name = BODY_FONT


# ============================================================ PART I: THE WORLD
part_divider(doc, "Part I", "The World",
             setting.get("tagline") or "Where the war is fought, and why.")

if prose.get("worldIntro"):
    heading(doc, "A World at War", 1)
    epigraph(doc, next_epigraph())
    body(doc, prose["worldIntro"])
    add_image(doc, "world_toll", 6.3)

heading(doc, "The Premise", 1)
if setting.get("premise"):
    body(doc, setting["premise"])
if setting.get("theTwist"):
    callout(doc, "The Twist", setting["theTwist"])
add_image(doc, "world_shatter", 6.3)

tl = prose.get("grandTimeline") or setting.get("timeline") or []
if tl:
    heading(doc, "Timeline", 1)
    add_table(doc, ["When", "What Happened"],
              [[b.get("when", ""), b.get("event", "")] for b in tl],
              widths=[Mm(38), Mm(132)], font_size=10)

terms = setting.get("terminology") or []
if terms:
    heading(doc, "The Tongue of the War", 2)
    body(doc, "Words you will hear across the battlefield:", space_after=3)
    add_table(doc, ["Term", "Meaning"],
              [[t.get("term", ""), t.get("meaning", "")] for t in terms],
              widths=[Mm(42), Mm(128)], font_size=10)

wm = setting.get("worldMechanics") or []
if wm:
    heading(doc, "Forces of the World", 1)
    body(doc, "Three forces shape every battle. They appear again as rules in Parts "
              "II and VI — the world is not just backdrop, it fights too.",
         space_after=4)
    for m in wm:
        heading(doc, m.get("name", "Force"), 2)
        body(doc, m.get("concept", ""), space_after=2)
        callout(doc, "On the table", m.get("ruleHook", ""))


# ============================================================ PART II: CORE RULES
part_divider(doc, "Part II", "Core Rules",
             "The whole engine. One number to roll. No spreadsheet.")

heading(doc, "The One Promise", 1)
epigraph(doc, next_epigraph())
body(doc, "There is one number you ever try to roll: a model's **target number**. "
          "Roll it or higher and you succeed. The entire game has only a handful of "
          "modifiers, each tied to something you can see on the table, and they "
          "never stack into a sum. If you can count to six, you can play.")
callout(doc, "The promise",
        "You will never add and subtract a chain of modifiers that cancel out. "
        "Every modifier changes a decision you make — or it isn't in the game.")

heading(doc, "Dice", 1)
body(doc, "You only need six-sided dice (**D6**). A natural **6 always succeeds** "
          "and a natural **1 always fails**. A few weapons roll a **D3** — roll "
          "a D6 and read it:")
add_table(doc, ["D6 result", "1 – 2", "3 – 4", "5 – 6"],
          [["D3 value", "1", "2", "3"]],
          widths=[Mm(40), Mm(35), Mm(35), Mm(35)])

heading(doc, "Measuring", 1)
body(doc, "Distances are in **inches**. Measure base-edge to base-edge. You may "
          "**measure anything at any time** — no guessing, no gotchas.")

heading(doc, "The Model Profile", 1)
add_table(doc, ["Stat", "Name", "What it does"],
          [["M", "Move", "Inches the model moves with one Move action."],
           ["SK", "Skill", "Target number to hit — used for shooting AND fighting."],
           ["DEF", "Defense", "Target number to save against a hit."],
           ["W", "Wounds", "Hits the model takes before it is removed (usually 1)."],
           ["A", "Attacks", "Dice the model rolls when it Fights in melee."]],
          widths=[Mm(18), Mm(30), Mm(122)])
body(doc, "Lower is better for **SK** and **DEF** (3+ beats 5+).", size=9.5, space_after=4)

heading(doc, "The Game Round", 1)
numbered(doc, "**Initiative.** Both players roll a D6; highest chooses who goes "
              "first this round (re-roll ties).")
numbered(doc, "**Alternate activations.** Take turns activating **one model each** "
              "until every model has activated.")
numbered(doc, "**End of round.** Resolve scenario and environment effects, then "
              "start the next round. A game is usually **4 rounds**.")

heading(doc, "Activations & Actions", 1)
body(doc, "An activated model takes **two actions**. Only **Move** may be taken "
          "twice; a model may **Shoot** at most once and **Fight** at most once.")
add_table(doc, ["Action", "Effect"],
          [["Move", "Move up to M inches."],
           ["Shoot", "Fire one weapon at a target in range and line of sight."],
           ["Fight", "Attack an enemy in base contact."],
           ["Aim", "Your next Shoot this activation ignores Cover and re-rolls one missed hit."],
           ["Rally", "Remove the Shaken condition from this model."]],
          widths=[Mm(24), Mm(146)])

heading(doc, "The Attack Sequence (shooting & fighting)", 1)
numbered(doc, "**Dice.** Shooting: roll the weapon's **Shots**. Fighting: roll the "
              "model's **A**, **+1 die if you Charged** (moved into base contact "
              "this activation).")
numbered(doc, "**Hit.** Each die that meets the attacker's **SK** is a hit.")
numbered(doc, "**Save.** The target rolls its **DEF** per hit; **+1 to saves if in "
              "Cover** (shooting only). Each failed save lets a hit through.")
numbered(doc, "**Damage.** Each unsaved hit removes the weapon's **Damage** in "
              "Wounds (usually 1; D3 for heavy weapons). At 0 Wounds, remove the model.")
callout(doc, "Line of sight & cover",
        "If you can draw a line to any part of the target, you can see it. A target "
        "**in cover** (inside terrain, or shot crosses terrain within 1\" of it) "
        "adds **+1 to its saves** vs shooting. Cover never stacks and never helps in "
        "melee. **Aim** and Sniper weapons ignore it.")

heading(doc, "Nerve (Morale)", 1)
body(doc, "A warband becomes **Broken** once it has lost **half or more** of its "
          "starting models. While Broken, at the start of each of your models' "
          "activations roll a D6: on a **1–2** that model is **Shaken** (Move "
          "actions only this activation); on 3+ it acts normally.")

ukw = council.get("universalKeywords") or []
if ukw:
    heading(doc, "Universal Keywords", 1)
    body(doc, "These special rules appear across many factions. Faction and unit "
              "entries reference them by name.", space_after=4)
    add_table(doc, ["Keyword", "Rule"],
              [[k.get("keyword", ""), k.get("text", "")] for k in ukw],
              widths=[Mm(40), Mm(130)], font_size=9.5)


# ============================================================ PART III: FACTIONS
part_divider(doc, "Part III", "The Factions",
             "Four powers. Four ways of waging war.")

dist_by_faction = {d.get("faction", ""): d for d in (council.get("distinctiveness") or [])}

for fidx, fac in enumerate(factions):
    doc.add_page_break()
    name = fac.get("name", "Faction")
    heading(doc, name, 1)
    if fac.get("battleCry"):
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(6)
        r = p.add_run('“' + fac["battleCry"] + '”')
        r.italic = True
        r.font.size = Pt(12)
        r.font.color.rgb = ACCENT
        r.font.name = BODY_FONT

    add_image(doc, "fac_%d" % fidx, 4.3)
    two_up_images(doc, "unit_%d_1" % fidx, "unit_%d_2" % fidx, {})

    d = dist_by_faction.get(name)
    if d:
        callout(doc, "How they play",
                "**" + d.get("playstyle", "") + "**  —  " + d.get("signatureMechanic", ""),
                fill=ZEBRA_FILL, accent_hex="2E4552")

    lore = fac.get("lore", {}) or {}
    lore_sections = [("Origin", "origin"), ("Culture", "culture"),
                     ("Why They Fight", "whyTheyFight"),
                     ("Leadership", "leadership"), ("Homeland", "homeland")]
    for label, key in lore_sections:
        if lore.get(key):
            heading(doc, label, 2)
            body(doc, lore[key])

    sig = fac.get("signatureRule") or {}
    if sig:
        callout(doc, "Signature Rule — " + sig.get("name", ""), sig.get("text", ""))

    fsr = fac.get("factionSpecialRules") or []
    if fsr:
        heading(doc, "Faction Special Rules", 2)
        for r in fsr:
            bullet(doc, "**" + r.get("name", "") + ":** " + r.get("text", ""))

    subs = fac.get("subfactions") or []
    if subs:
        heading(doc, "Subfactions", 2)
        for s in subs:
            body(doc, "**" + s.get("name", "") + "** — " + s.get("twist", ""),
                 space_after=2)
            callout(doc, s.get("ruleName", "Rule"), s.get("ruleText", ""),
                    fill=BOX_FILL, accent_hex="8A6D1F")

    units = fac.get("units") or []
    if units:
        heading(doc, "Army List", 2)
        rows = []
        for u in units:
            wpns = ", ".join(u.get("weapons", []) or [])
            extra = (u.get("wargear") or [])
            if extra:
                wpns = (wpns + "; " + ", ".join(extra)).strip("; ")
            rules = ", ".join(u.get("specialRules", []) or []) or "—"
            rows.append([
                u.get("name", ""), str(u.get("M", "—")), str(u.get("SK", "—")),
                str(u.get("DEF", "—")), str(u.get("W", "—")),
                str(u.get("A", "—")), wpns or "—", rules,
                str(u.get("points", "—")),
            ])
        add_table(doc,
                  ["Unit", "M", "SK", "DEF", "W", "A", "Weapons & Wargear", "Special Rules", "Pts"],
                  rows,
                  widths=[Mm(28), Mm(10), Mm(11), Mm(12), Mm(8), Mm(8), Mm(42), Mm(33), Mm(11)],
                  font_size=8.5)

        # notable unit flavor
        flav = [u for u in units if u.get("flavor")]
        if flav:
            heading(doc, "Notable Units", 2)
            for u in flav:
                bullet(doc, "**" + u.get("name", "") + "** — " + u.get("flavor", ""),
                       size=10)

    chars = fac.get("characters") or []
    if chars:
        heading(doc, "Heroes & Characters", 2)
        for ch in chars:
            character_box(doc, ch)

    wt = fac.get("weaponsTable") or []
    if wt:
        heading(doc, "Faction Wargear", 2)
        add_table(doc, ["Weapon", "Range", "Shots", "Damage", "Special"],
                  [[w.get("name", ""), w.get("range", ""), w.get("shots", ""),
                    w.get("damage", ""), w.get("special", "")] for w in wt],
                  widths=[Mm(34), Mm(20), Mm(18), Mm(20), Mm(78)], font_size=9)

# rivalries
riv = council.get("rivalries") or []
if riv:
    doc.add_page_break()
    heading(doc, "Rivalries & Grudges", 1)
    body(doc, "Old hatreds shape the war. Use these when picking scenarios and "
              "campaign matchups.", space_after=4)
    add_table(doc, ["Between", "The Grudge"],
              [[r.get("between", ""), r.get("nature", "")] for r in riv],
              widths=[Mm(48), Mm(122)], font_size=10)


# ============================================================ PART IV: SCENARIOS
part_divider(doc, "Part IV", "Scenarios",
             "Reasons to fight, and ways to win that aren't just killing.")

for i, sc in enumerate(scenarios, 1):
    if i > 1:
        doc.add_paragraph()
    heading(doc, f"{i}.  {sc.get('name', 'Scenario')}", 1)
    if sc.get("hook"):
        body(doc, sc["hook"], italic=True, color=STEEL, space_after=4)
    add_image_path(doc, os.path.join(ASSETS, "map_%d.png" % (i - 1)), 5.7,
                   caption="Suggested deployment — " + sc.get("name", ""))
    for label, key in [("Setup", "setup"), ("Deployment", "deployment"),
                       ("Objectives", "objectives"), ("Twist", "twist")]:
        if sc.get(key):
            heading(doc, label, 2)
            body(doc, sc[key])
    srs = sc.get("specialRules") or []
    if srs:
        heading(doc, "Special Rules", 2)
        for r in srs:
            bullet(doc, "**" + r.get("name", "") + ":** " + r.get("text", ""))
    meta_bits = []
    if sc.get("length"):
        meta_bits.append("**Length:** " + str(sc["length"]))
    if sc.get("victory"):
        meta_bits.append("**Victory:** " + sc["victory"])
    if meta_bits:
        callout(doc, "Winning", "   —   ".join(meta_bits))


# ============================================================ PART V: CAMPAIGN
if campaign:
    part_divider(doc, "Part V", "Campaign",
                 campaign.get("name", "The Long War"))
    heading(doc, campaign.get("name", "Campaign"), 1)
    epigraph(doc, next_epigraph())
    if campaign.get("overview"):
        body(doc, campaign["overview"])
    if campaign.get("howItWorks"):
        heading(doc, "How It Works", 2)
        body(doc, campaign["howItWorks"])
    prog = campaign.get("progression") or []
    if prog:
        heading(doc, "Advancement", 2)
        for r in prog:
            bullet(doc, "**" + r.get("name", "") + ":** " + r.get("text", ""))
    if campaign.get("territories"):
        heading(doc, "Territory", 2)
        body(doc, campaign["territories"])
    inj = campaign.get("injuries") or []
    if inj:
        heading(doc, "Battle Scars", 2)
        body(doc, "Roll for any model taken out of action that survives the war.",
             space_after=3)
        add_table(doc, ["D6", "Result"],
                  [[r.get("roll", ""), r.get("result", "")] for r in inj],
                  widths=[Mm(20), Mm(150)], font_size=10)


# ============================================================ PART VI: ADVANCED
if advanced:
    part_divider(doc, "Part VI", "Advanced & Optional Rules",
                 "Let the world fight back.")
    tt = advanced.get("terrainTypes") or []
    if tt:
        heading(doc, "Terrain Types", 1)
        for r in tt:
            bullet(doc, "**" + r.get("name", "") + ":** " + r.get("text", ""))
    ev = advanced.get("environmentEvents") or []
    if ev:
        heading(doc, "Environment Events", 1)
        body(doc, "At the start of each round (optional), roll a D6 and apply the "
                  "event. The world is never just scenery.", space_after=3)
        add_table(doc, ["D6", "Event", "Effect"],
                  [[e.get("roll", ""), e.get("name", ""), e.get("text", "")] for e in ev],
                  widths=[Mm(16), Mm(40), Mm(114)], font_size=9.5)
    opt = advanced.get("optionalRules") or []
    if opt:
        heading(doc, "Optional Rules", 1)
        for r in opt:
            bullet(doc, "**" + r.get("name", "") + ":** " + r.get("text", ""))


# ============================================================ PART VII: QUICK REF
part_divider(doc, "Part VII", "Quick Reference", "Print this. Keep it between you.")

heading(doc, "The Round", 2)
body(doc, "**1.** Roll Initiative (high picks who goes first).  **2.** Alternate "
          "activating one model each.  **3.** End of round: scenario + environment, "
          "then repeat. Game = 4 rounds.", space_after=4)
heading(doc, "Activation = 2 actions", 2)
body(doc, "Move (repeatable) · Shoot (once) · Fight (once) · Aim · Rally.",
     space_after=4)
heading(doc, "Attack sequence", 2)
add_table(doc, ["Step", "Do this"],
          [["1. Dice", "Shoot: roll Shots. Fight: roll A (+1 if you Charged)."],
           ["2. Hit", "Each die ≥ SK = hit. Nat 6 always hits, 1 always misses."],
           ["3. Save", "Target rolls DEF per hit; +1 to saves in Cover (shooting only)."],
           ["4. Damage", "Each unsaved hit removes Damage Wounds. 0 = removed."]],
          widths=[Mm(26), Mm(144)], font_size=9.5)
heading(doc, "The complete modifier list", 2)
add_table(doc, ["Modifier", "Effect", "When"],
          [["Cover", "+1 to saves", "Target in/behind terrain (shooting only)"],
           ["Charge", "+1 attack die", "Moved into base contact this activation"],
           ["Aim", "Ignore cover, re-roll 1 miss", "Spent the Aim action"]],
          widths=[Mm(24), Mm(54), Mm(92)], font_size=9.5)
body(doc, "Everything else is a named keyword that says exactly what it does. No "
          "arithmetic, ever.", size=9.5, italic=True, space_after=4)

dist = council.get("distinctiveness") or []
if dist:
    heading(doc, "Faction signatures", 2)
    add_table(doc, ["Faction", "Plays like", "Signature"],
              [[d.get("faction", ""), d.get("playstyle", ""), d.get("signatureMechanic", "")] for d in dist],
              widths=[Mm(34), Mm(46), Mm(90)], font_size=9)

heading(doc, "D3 from a D6", 2)
add_table(doc, ["D6", "1–2", "3–4", "5–6"],
          [["D3", "1", "2", "3"]],
          widths=[Mm(34), Mm(30), Mm(30), Mm(30)], font_size=9.5)


# ============================================================ CLOSING FICTION
if prose.get("closingFiction"):
    doc.add_page_break()
    for _ in range(3):
        doc.add_paragraph()
    fiction(doc, prose["closingFiction"])

doc.save(os.path.join(HERE, "sixes_rulebook.docx"))
print("Wrote sixes_rulebook.docx")
print(f"  world={WORLD!r} factions={len(factions)} scenarios={len(scenarios)} "
      f"keywords={len(ukw)} rivalries={len(riv)}")
