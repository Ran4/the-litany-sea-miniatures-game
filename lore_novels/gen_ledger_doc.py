# /// script
# requires-python = ">=3.11"
# dependencies = ["python-docx"]
# ///
"""Lay out the_ledger.md as a book-styled .docx, convert to .pdf, and open it."""
import re, subprocess, sys
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.enum.section import WD_SECTION
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

HERE = Path(__file__).parent
SRC = HERE / "the_ledger.md"
DOCX = HERE / "the_ledger.docx"
BODY_FONT = "EB Garamond"
INK = RGBColor(0x1A, 0x1A, 0x1A)


def smart(t: str) -> str:
    """Straight quotes -> typographic quotes (text already uses real em dashes)."""
    t = re.sub(r"(\w)'", "\\1’", t)        # apostrophe / closing single
    t = t.replace("'", "‘")                  # opening single
    out, openq = [], True
    for ch in t:
        if ch == '"':
            out.append("“" if openq else "”")
            openq = not openq
        else:
            out.append(ch)
    return "".join(out)


def add_page_number(paragraph):
    run = paragraph.add_run()
    for kind, txt in (("begin", None), (None, "PAGE"), ("end", None)):
        if kind:
            el = OxmlElement("w:fldChar")
            el.set(qn("w:fldCharType"), kind)
        else:
            el = OxmlElement("w:instrText")
            el.set(qn("xml:space"), "preserve")
            el.text = txt
        run._r.append(el)


# ---- parse the manuscript ----
title = subtitle = epigraph = ""
blocks = []  # ("chapter", text) | ("para", text)
for raw in SRC.read_text(encoding="utf-8").splitlines():
    line = raw.rstrip()
    if line.startswith("TITLE:"):
        title = line[6:].strip()
    elif line.startswith("SUBTITLE:"):
        subtitle = line[9:].strip()
    elif line.startswith("EPIGRAPH:"):
        epigraph = line[9:].strip()
    elif line.startswith("## "):
        blocks.append(("chapter", line[3:].strip()))
    elif line.strip():
        blocks.append(("para", line.strip()))

# ---- build the document ----
doc = Document()
sec = doc.sections[0]
sec.page_height, sec.page_width = Cm(29.7), Cm(21.0)  # A4
sec.top_margin = sec.bottom_margin = Cm(2.6)
sec.left_margin = sec.right_margin = Cm(3.0)

normal = doc.styles["Normal"]
normal.font.name = BODY_FONT
normal.font.size = Pt(12)
normal.font.color.rgb = INK

body = doc.styles.add_style("Body", 1)  # WD_STYLE_TYPE.PARAGRAPH
body.base_style = normal
bpf = body.paragraph_format
bpf.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
bpf.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
bpf.line_spacing = 1.4
bpf.first_line_indent = Cm(0.6)
bpf.space_after = Pt(0)

# page numbers, centered in footer
pn = sec.footer.paragraphs[0]
pn.alignment = WD_ALIGN_PARAGRAPH.CENTER
pn.style = normal
add_page_number(pn)

# ---- title page ----
for _ in range(4):
    doc.add_paragraph()
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run(smart(title))
r.font.name = BODY_FONT
r.font.size = Pt(40)
r.font.small_caps = True
r.font.color.rgb = INK

if subtitle:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(10)
    r = p.add_run(smart(subtitle))
    r.italic = True
    r.font.size = Pt(15)

if epigraph:
    for _ in range(2):
        doc.add_paragraph()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(smart(epigraph))
    r.italic = True
    r.font.size = Pt(12.5)

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_before = Pt(36)
r = p.add_run("— The Litany Sea —")
r.font.size = Pt(11)
r.font.color.rgb = RGBColor(0x70, 0x70, 0x70)

# ---- body ----
first_after_heading = False
for kind, text in blocks:
    if kind == "chapter":
        h = doc.add_paragraph()
        h.alignment = WD_ALIGN_PARAGRAPH.CENTER
        h.paragraph_format.page_break_before = True
        h.paragraph_format.space_after = Pt(20)
        h.paragraph_format.keep_with_next = True
        r = h.add_run(smart(text))
        r.bold = True
        r.font.name = BODY_FONT
        r.font.size = Pt(17)
        r.font.small_caps = True
        r.font.color.rgb = INK
        first_after_heading = True
    else:
        para = doc.add_paragraph(style=body)
        if first_after_heading:
            para.paragraph_format.first_line_indent = Cm(0)
            first_after_heading = False
        para.add_run(smart(text))

doc.save(DOCX)
words = sum(len(t.split()) for k, t in blocks if k == "para")
print(f"wrote {DOCX.name}  ({len([1 for k,_ in blocks if k=='chapter'])} sections, ~{words} words)")

# ---- convert to PDF + open ----
subprocess.run(
    ["soffice", "--headless", "--convert-to", "pdf", "--outdir", str(HERE), str(DOCX)],
    check=True,
)
pdf = DOCX.with_suffix(".pdf")
print(f"wrote {pdf.name}")
subprocess.Popen(["xdg-open", str(pdf)], stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)
